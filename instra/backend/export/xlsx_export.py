"""XRF 数据 xlsx 导出：表头 + 单位行 + 多样品并排数据。

定量与 UQ 同布局：全量输出不看 use_report；UQ 可按元素/氧化物口径导出。
每个元素/氧化物对可选 % / ppm / ppb，数值按本次导出的有效数字舍入。
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from domain import conversion, format as value_format, ordering

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
_UNIT_FILL = PatternFill("solid", fgColor="F5F7FA")


def _sample_label(row):
    """首列样品名：关联 LIMS 样品优先，否则原始样品名。"""
    return row.get("lims_name") or row.get("sample_name") or ""


def _value_with_key(name, value, elements, oxides):
    pair = conversion.composition_pairs(
        [{"name": name, "value": value}], elements, oxides)[0]
    return {"value": value, "key": conversion.composition_pair_key(pair)}


def _fetch_quant_rows(db, ids, elements, oxides):
    marks = ",".join("?" for _ in ids)
    rows = db.execute(f"""SELECT xa.id,xa.external_id,xa.sample_name,xa.method,xa.batch,
        xa.analyzed_at,s.name AS lims_name
        FROM xrf_analyses xa LEFT JOIN samples s ON s.id=xa.sample_id
        WHERE xa.id IN ({marks}) ORDER BY xa.analyzed_at NULLS LAST, xa.id""", ids).fetchall()
    result = []
    for row in rows:
        values = {}
        for value in db.execute(
                "SELECT name,value FROM xrf_values WHERE analysis_id=?", (row["id"],)):
            percent = float(value["value"])
            values[value["name"]] = _value_with_key(
                value["name"], percent, elements, oxides)
        result.append((dict(row), values))
    return result


def _uq_values_for_basis(db, analysis_id, basis, elements, oxides):
    """UQ 一条分析（xrf_analyses.id）在指定口径下的 {显示名: 含量}。

    basis 与数据原生口径不一致时按 common_oxides 系数换算；未知名字原样保留。
    """
    values = db.execute("SELECT name,value,alt_name FROM xrf_values WHERE analysis_id=?",
                        (analysis_id,)).fetchall()
    out = {}
    for value in values:
        name, val = value["name"], float(value["value"])
        display_name, display_value, _ = conversion.to_basis(
            name, value["alt_name"] or "", val, basis, elements, oxides)
        pair = conversion.composition_pairs(
            [{"name": name, "value": val}], elements, oxides)[0]
        out[display_name] = {
            "value": display_value,
            "key": conversion.composition_pair_key(pair),
        }
    return out


def _write_sheet(title, header_names, rows, units, significant_digits):
    """rows: [(样品名, {项名: {value,key}})]；返回 Workbook 字节。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["样品", *header_names])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    header_keys = {}
    for name in header_names:
        header_keys[name] = next(
            (values[name]["key"] for _, values in rows if name in values), name)
    resolved_units = {
        name: units.get(header_keys[name]) or value_format.default_xrf_unit(max(
            (abs(values[name]["value"]) for _, values in rows if name in values),
            default=0,
        ))
        for name in header_names
    }
    ws.append(["单位", *[resolved_units[name] for name in header_names]])
    for cell in ws[2]:
        cell.fill = _UNIT_FILL
        cell.font = Font(italic=True, color="667085")
    for label, values in rows:
        exported = []
        for name in header_names:
            item = values.get(name)
            if item is None:
                exported.append(None)
                continue
            converted = value_format.convert_percent(item["value"], resolved_units[name])
            exported.append(value_format.round_significant(converted, significant_digits))
        ws.append([label, *exported])
    ws.column_dimensions["A"].width = 22
    for index in range(2, len(header_names) + 2):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = 12
    ws.freeze_panes = "B3"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def load_unit_fields(db, ids, kind, order_template_id=None, order_mode="template"):
    """返回所选扫描的元素/氧化物对及其自动单位，供导出抽屉使用。"""
    if not ids:
        return []
    elements, oxides = conversion.load_reference(db)
    marks = ",".join("?" for _ in ids)
    rows = db.execute(f"""SELECT xv.name,xv.value FROM xrf_values xv
        JOIN xrf_analyses xa ON xa.id=xv.analysis_id
        WHERE xa.id IN ({marks}) AND COALESCE(xa.kind,'quant')=?""",
        [*ids, kind]).fetchall()
    fields = {}
    for row in rows:
        pair = conversion.composition_pairs(
            [{"name": row["name"], "value": float(row["value"])}], elements, oxides)[0]
        key = conversion.composition_pair_key(pair)
        maximum = max(abs(float(value)) for value in
                      (pair["oxide_value"], pair["element_value"]) if value is not None)
        if key not in fields:
            fields[key] = {
                "key": key,
                "label": conversion.composition_pair_label(pair),
                "oxide_name": pair["oxide_name"],
                "element_name": pair["element_name"],
                "max_percent": maximum,
                "max_element_percent": abs(float(
                    pair["element_value"] if pair["element_value"] is not None
                    else pair["oxide_value"])),
            }
        else:
            fields[key]["max_percent"] = max(fields[key]["max_percent"], maximum)
            element_value = (pair["element_value"] if pair["element_value"] is not None
                             else pair["oxide_value"])
            fields[key]["max_element_percent"] = max(
                fields[key]["max_element_percent"], abs(float(element_value)))

    order_list = ordering.load_order_list(db, order_template_id)
    rank = {conversion.normalize(name): i for i, name in enumerate(order_list)}

    def field_rank(field):
        names = [conversion.normalize(name) for name in
                 (field["oxide_name"], field["element_name"]) if name]
        matched = [rank[name] for name in names if name in rank]
        return (min(matched) if matched else len(rank), field["label"].casefold())

    result = _sort_fields(list(fields.values()), order_mode, field_rank)
    for field in result:
        field["default_unit"] = value_format.default_xrf_unit(field.pop("max_percent"))
        field.pop("max_element_percent")
    return result


def _sort_fields(fields, order_mode, template_rank=None):
    """按跨样品最大元素含量或模板 rank 排字段。"""
    if order_mode == "content":
        return sorted(fields, key=lambda field: (-field["max_element_percent"],
                                                  field["label"].casefold()))
    return sorted(fields, key=template_rank) if template_rank else fields


def _resolved_units(fields, units):
    resolved = {field["key"]: field["default_unit"] for field in fields}
    resolved.update(units or {})
    return resolved


def _ordered_names(rows, order_list, fields, order_mode):
    names = ordering.order_items([name for _, values in rows for name in values],
                                 order_list if order_mode == "template" else [])
    if order_mode != "content":
        return names
    pair_rank = {field["key"]: index for index, field in enumerate(fields)}
    name_keys = {name: next(values[name]["key"] for _, values in rows if name in values)
                 for name in names}
    return sorted(names, key=lambda name: (pair_rank.get(name_keys[name], len(pair_rank)),
                                           name.casefold()))


def build_quant_xlsx(db, ids, order_template_id=None, units=None, significant_digits=5,
                     order_mode="template"):
    elements, oxides = conversion.load_reference(db)
    rows = _fetch_quant_rows(db, ids, elements, oxides)
    if not rows:
        raise ValueError("没有可导出的定量扫描")
    order_list = ordering.load_order_list(db, order_template_id)
    fields = load_unit_fields(db, ids, "quant", order_template_id, order_mode)
    names = _ordered_names(rows, order_list, fields, order_mode)
    resolved_units = _resolved_units(fields, units)
    return _write_sheet("XRF定量", names,
                        [(_sample_label(row), values) for row, values in rows],
                        resolved_units, significant_digits)


def build_uq_xlsx(db, ids, basis, order_template_id=None, units=None,
                  significant_digits=5, order_mode="content"):
    """ids 为 xrf_analyses.id（kind='uq' 的扫描）。"""
    elements, oxides = conversion.load_reference(db)
    order_list = ordering.load_order_list(db, order_template_id)
    rows = []
    for aid in ids:
        meta = db.execute("""SELECT xa.id,xa.sample_name,s.name AS lims_name
            FROM xrf_analyses xa LEFT JOIN samples s ON s.id=xa.sample_id
            WHERE xa.id=? AND COALESCE(xa.kind,'quant')='uq'""", (aid,)).fetchone()
        if meta is None:
            continue
        values = _uq_values_for_basis(db, aid, basis, elements, oxides)
        rows.append((dict(meta), values))
    if not rows:
        raise ValueError("没有可导出的 UniQuant 分析")
    fields = load_unit_fields(db, ids, "uq", order_template_id, order_mode)
    names = _ordered_names(rows, order_list, fields, order_mode)
    resolved_units = _resolved_units(fields, units)
    return _write_sheet("UniQuant", names,
                        [(_sample_label(row), values) for row, values in rows],
                        resolved_units, significant_digits)


def quant_filename():
    return f"xrf-quant-{datetime.now():%Y%m%d-%H%M%S}.xlsx"


def uq_filename(basis):
    label = "元素" if basis == "element" else "氧化物"
    return f"xrf-uq-{label}-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
