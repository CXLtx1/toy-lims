import io
import unittest

from openpyxl import load_workbook

from export.xlsx_export import _sort_fields, _write_sheet


class TestXlsxUnits(unittest.TestCase):
    def test_content_order_uses_largest_value_from_any_sample(self):
        fields = [
            {"label": "Cl", "max_element_percent": 20},
            {"label": "Al", "max_element_percent": 40},
            {"label": "Cu", "max_element_percent": 30},
        ]
        self.assertEqual([field["label"] for field in
                          _sort_fields(fields, "content")], ["Al", "Cu", "Cl"])

    def test_unit_row_conversion_and_significant_digits(self):
        data = _write_sheet("test", ["CaO", "Cl"], [
            ("sample-a", {
                "CaO": {"value": 0.2, "key": "cao|ca"},
                "Cl": {"value": 0.000000123456, "key": "cl|cl"},
            }),
        ], {"cao|ca": "ppm", "cl|cl": "ppb"}, 4)
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual([cell.value for cell in ws[1]], ["样品", "CaO", "Cl"])
        self.assertEqual([cell.value for cell in ws[2]], ["单位", "ppm", "ppb"])
        self.assertEqual([cell.value for cell in ws[3]], ["sample-a", 2000, 1.235])
        self.assertEqual(ws.freeze_panes, "B3")


if __name__ == "__main__":
    unittest.main()
