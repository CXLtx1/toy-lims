import io
import inspect
import unittest

from openpyxl import load_workbook

import app as lims
from client_helpers import browser_client
from postgres_case import PostgresTestCase


class BusinessExcelApiTest(PostgresTestCase):
    def setUp(self):
        self.provision_database(lims)
        lims.app.config.update(TESTING=True, AUTH_DISABLED=True)
        self.client = browser_client(self, lims.app)
        self.meta = self.client.get("/api/meta").get_json()

    def create_sample(self, name="Excel业务样品", xrf=False):
        aid = next(a["id"] for a in self.meta["analytes"] if a["name"] == "Ag")
        instrument = next(i for i in self.meta["instruments"] if aid in i["analytes"] and i["itype"] == "ppm")
        dilution = self.meta["dilutions"][0]
        payload = {"name": name, "category": "业务测试", "xrf": int(xrf), "preps": [{
            "name": name + "-溶液", "mass_g": 1.25, "volume_ml": 100,
            "dilution_id": dilution["id"], "analyte_ids": [aid],
            "instrument_map": {str(aid): {"instrument_id": instrument["id"]}},
        }]}
        if xrf:
            payload["xrf_method_id"] = next(m["id"] for m in self.meta["methods"] if m["itype"] == "xrf")
        response = self.client.post("/api/samples", json=payload)
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        sid = response.get_json()["id"]
        self.client.put(f"/api/samples/{sid}/status", json={"status": "queued"})
        self.client.put(f"/api/samples/{sid}/status", json={"status": "measuring"})
        return sid

    def download(self, path):
        response = self.client.get(path)
        if response.status_code != 200:
            self.fail(f"{path} returned {response.status_code}: {response.get_data(as_text=True)}")
        self.assertIn("application/vnd.openxmlformats", response.content_type)
        return response.data

    @staticmethod
    def upload_data(blob, name="business.xlsx"):
        return {"file": (io.BytesIO(blob), name)}

    def test_overview_uses_inclusive_creation_dates_and_filters(self):
        first, second = self.create_sample("首日样品"), self.create_sample("次日样品")
        db = self.connect()
        try:
            db.execute("UPDATE samples SET created_at='2026-08-01 00:00:00' WHERE id=%s", (first,))
            db.execute("UPDATE samples SET created_at='2026-08-02 23:59:59' WHERE id=%s", (second,))
            db.commit()
        finally:
            db.close()
        blob = self.download("/api/excel/samples-overview?date_from=2026-08-01&date_to=2026-08-02")
        wb = load_workbook(io.BytesIO(blob), read_only=True)
        names = [row[1] for row in wb["样品总览"].iter_rows(min_row=4, values_only=True)]
        self.assertEqual(["首日样品", "次日样品"], names)
        self.assertEqual("toy-lims-business-workbook", wb["说明"]["B3"].value)
        wb.close()

    def test_plan_export_create_round_trip_and_overwrite_preserves_task(self):
        sid = self.create_sample()
        original = self.client.get(f"/api/samples/{sid}").get_json()
        task_id = original["items"][0]["id"]
        blob = self.download(f"/api/excel/samples/{sid}/detail")
        created = self.client.post("/api/excel/samples/create", data=self.upload_data(blob),
                                   content_type="multipart/form-data")
        self.assertEqual(200, created.status_code, created.get_data(as_text=True))
        self.assertNotEqual(sid, created.get_json()["id"])
        self.assertNotEqual(original["sample"]["lims_no"], created.get_json()["lims_no"])

        wb = load_workbook(io.BytesIO(blob))
        for row in wb["样品信息"].iter_rows(min_row=4):
            if row[0].value == "样品名称":
                row[1].value = "Excel覆盖名称"
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        overwritten = self.client.post(f"/api/excel/samples/{sid}/detail",
                                       data={"file": (io.BytesIO(changed.getvalue()), "plan.xlsx")},
                                       content_type="multipart/form-data")
        self.assertEqual(200, overwritten.status_code, overwritten.get_data(as_text=True))
        detail = self.client.get(f"/api/samples/{sid}").get_json()
        self.assertEqual("Excel覆盖名称", detail["sample"]["name"])
        self.assertEqual(task_id, detail["items"][0]["id"])

    def test_plan_round_trip_preserves_multistage_dilution(self):
        aid = next(a["id"] for a in self.meta["analytes"] if a["name"] == "Ag")
        instrument = next(i for i in self.meta["instruments"] if aid in i["analytes"] and i["itype"] == "ppm")
        dilution = next(d for d in self.meta["dilutions"] if d["label"] == "5/250")
        response = self.client.post("/api/samples", json={"name": "Excel多级稀释", "preps": [{
            "name": "Excel多级稀释*2500", "mass_g": 1, "volume_ml": 250,
            "dilution_ids": [dilution["id"], dilution["id"]], "analyte_ids": [aid],
            "instrument_map": {str(aid): {"instrument_id": instrument["id"]}},
        }]})
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        blob = self.download(f"/api/excel/samples/{response.get_json()['id']}/detail")
        workbook = load_workbook(io.BytesIO(blob))
        prep_row = next(workbook["溶样方案"].iter_rows(min_row=4, values_only=True))
        self.assertEqual(f"{dilution['id']},{dilution['id']}", prep_row[4])
        self.assertEqual("5/250 × 5/250", prep_row[5])
        workbook.close()

        created = self.client.post("/api/excel/samples/create", data=self.upload_data(blob),
                                   content_type="multipart/form-data")
        self.assertEqual(200, created.status_code, created.get_data(as_text=True))
        prep = self.client.get(f"/api/samples/{created.get_json()['id']}").get_json()["preps"][0]
        self.assertEqual([dilution["id"], dilution["id"]], prep["dilution_ids"])
        self.assertEqual(2500, prep["factor"])

    def test_plan_round_trip_preserves_formula_like_text(self):
        sid = self.create_sample("-Excel安全文本")
        blob = self.download(f"/api/excel/samples/{sid}/detail")
        wb = load_workbook(io.BytesIO(blob))
        name_cell = next(row[1] for row in wb["样品信息"].iter_rows(min_row=4)
                         if row[0].value == "样品名称")
        self.assertEqual("-Excel安全文本", name_cell.value)
        self.assertTrue(name_cell.quotePrefix)
        wb.close()
        response = self.client.post(f"/api/excel/samples/{sid}/detail",
                                    data=self.upload_data(blob), content_type="multipart/form-data")
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        self.assertEqual("-Excel安全文本", self.client.get(f"/api/samples/{sid}").get_json()["sample"]["name"])

    def test_direct_xrf_task_plan_round_trip_and_delete_are_protected(self):
        sid = self.create_sample(xrf=True)
        detail = self.client.get(f"/api/samples/{sid}").get_json()
        aid = detail["items"][0]["analyte_id"]
        xrf_instrument = next(i for i in self.meta["instruments"] if i["itype"] == "xrf")
        xrf_method = next(m for m in self.meta["methods"] if m["itype"] == "xrf")
        db = self.connect()
        try:
            xrf_task_id = db.execute("""INSERT INTO sample_analytes(sample_id,preparation_id,analyte_id,instrument_id,method_id)
                          VALUES(%s,NULL,%s,%s,%s) RETURNING id""",
                          (sid, aid, xrf_instrument["id"], xrf_method["id"])).fetchone()[0]
            db.commit()
        finally:
            db.close()
        blob = self.download(f"/api/excel/samples/{sid}/detail")
        unchanged = self.client.post(f"/api/excel/samples/{sid}/detail",
                                     data=self.upload_data(blob), content_type="multipart/form-data")
        self.assertEqual(200, unchanged.status_code, unchanged.get_data(as_text=True))
        wb = load_workbook(io.BytesIO(self.download(f"/api/excel/samples/{sid}/detail")))
        ws = wb["检测方案"]
        for row in range(ws.max_row, 3, -1):
            if ws.cell(row, 1).value == xrf_task_id:
                ws.delete_rows(row)
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        rejected = self.client.post(f"/api/excel/samples/{sid}/detail",
                                    data={"file": (io.BytesIO(changed.getvalue()), "delete-xrf.xlsx")},
                                    content_type="multipart/form-data")
        self.assertEqual(400, rejected.status_code)
        self.assertIn("XRF仪器任务", rejected.get_json()["error"])
        db = self.connect()
        try:
            self.assertIsNotNone(db.execute("SELECT id FROM sample_analytes WHERE id=%s", (xrf_task_id,)).fetchone())
        finally:
            db.close()

    def test_regular_data_overwrite_aux_and_xrf_unchanged(self):
        sid = self.create_sample(xrf=True)
        detail = self.client.get(f"/api/samples/{sid}").get_json()
        task_id = detail["items"][0]["id"]
        db = self.connect()
        try:
            analysis_id = db.execute(
                "INSERT INTO xrf_analyses(sample_id,external_id,method) VALUES(%s,%s,%s) RETURNING id",
                (sid, "X-1", "UQ")).fetchone()[0]
            db.execute("INSERT INTO xrf_values(analysis_id,name,value,use_report) VALUES(%s,%s,%s,1)", (analysis_id, "Ag", 9.9))
            db.commit()
        finally:
            db.close()
        blob = self.download(f"/api/excel/samples/{sid}/data")
        wb = load_workbook(io.BytesIO(blob))
        ws = wb["数据录入"]
        self.assertEqual(task_id, ws["A4"].value)
        ws["H4"] = 12.345
        ws["I4"] = 1
        ws["K4"] = '{"V":12.345}'
        ws["L4"] = '{"use":true,"expected":10,"measured":9.8}'
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        response = self.client.post(f"/api/excel/samples/{sid}/data",
                                    data={"file": (io.BytesIO(changed.getvalue()), "data.xlsx")},
                                    content_type="multipart/form-data")
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        updated = self.client.get(f"/api/samples/{sid}").get_json()["items"][0]
        self.assertEqual(12.345, updated["readings"][0]["raw"])
        self.assertEqual('{"use": true, "expected": 10, "measured": 9.8}', updated["aux"])
        xrf = self.client.get(f"/api/xrf/samples/{sid}").get_json()
        self.assertEqual(9.9, xrf["analyses"][0]["values"][0]["value"])

    def test_formula_and_wrong_sample_are_rejected(self):
        sid, other = self.create_sample("公式样品"), self.create_sample("其他样品")
        blob = self.download(f"/api/excel/samples/{sid}/data")
        plan_blob = self.download(f"/api/excel/samples/{sid}/detail")
        profile = self.client.post(f"/api/excel/samples/{sid}/data", data=self.upload_data(plan_blob),
                                   content_type="multipart/form-data")
        self.assertEqual(400, profile.status_code)
        self.assertIn("业务类型错误", profile.get_json()["error"])
        wb = load_workbook(io.BytesIO(blob))
        wb["数据录入"]["H4"] = "=1+1"
        invalid = io.BytesIO()
        wb.save(invalid)
        wb.close()
        formula = self.client.post(f"/api/excel/samples/{sid}/data",
                                   data={"file": (io.BytesIO(invalid.getvalue()), "formula.xlsx")},
                                   content_type="multipart/form-data")
        self.assertEqual(400, formula.status_code)
        self.assertIn("不允许使用公式", formula.get_json()["error"])
        wrong = self.client.post(f"/api/excel/samples/{other}/data", data=self.upload_data(blob),
                                 content_type="multipart/form-data")
        self.assertEqual(409, wrong.status_code)

    def test_report_uses_json_payload_values_and_expected_sheets(self):
        sid = self.create_sample("报告样品")
        task = self.client.get(f"/api/samples/{sid}").get_json()["items"][0]
        reading = self.client.post("/api/readings", json={"sample_analyte_id": task["id"]}).get_json()["id"]
        self.client.put(f"/api/readings/{reading}", json={"raw": 10, "extra": {}, "use_avg": True})
        payload = self.client.get(f"/api/report/{sid}").get_json()
        blob = self.download(f"/api/excel/reports/{sid}")
        wb = load_workbook(io.BytesIO(blob), data_only=True)
        self.assertEqual(["说明", "分析报告", "原始分析记录", "计算明细"], wb.sheetnames)
        self.assertEqual(payload["groups"][0]["final"]["value"], wb["分析报告"]["C13"].value)
        self.assertIn("草稿", wb["说明"]["A1"].value)
        wb.close()

    def test_manual_report_override_skips_unchanged_and_updates_excel(self):
        sid = self.create_sample("手工报告样品")
        task = self.client.get(f"/api/samples/{sid}").get_json()["items"][0]
        reading = self.client.post("/api/readings", json={"sample_analyte_id": task["id"]}).get_json()["id"]
        self.client.put(f"/api/readings/{reading}", json={"raw": 10, "extra": {}, "use_avg": True})
        payload = self.client.get(f"/api/report/{sid}").get_json()
        defaults = payload["default_report_rows"]

        unchanged = self.client.put(f"/api/reports/{sid}/manual", json={
            "rows": defaults, "reason": "核对系统结果",
        })
        self.assertFalse(unchanged.get_json()["changed"])
        self.assertIsNone(self.client.get(f"/api/report/{sid}").get_json()["manual_report"])

        rows = [dict(defaults[0]), {
            "item": "补充说明项", "result": "合格", "unit": "", "note": "人工增加", "include": True,
        }]
        rows[0]["result"] = "12.34"
        changed = self.client.put(f"/api/reports/{sid}/manual", json={
            "rows": rows, "reason": "补录纸质检测结果",
        })
        self.assertTrue(changed.get_json()["changed"])
        manual_payload = self.client.get(f"/api/report/{sid}").get_json()
        self.assertEqual(rows, manual_payload["report_rows"])
        excluded = self.client.put(f"/api/samples/{sid}/report-print", json={
            "excludes": ["m:补充说明项"],
        })
        self.assertEqual(200, excluded.status_code, excluded.get_data(as_text=True))
        self.assertFalse(self.client.get(f"/api/report/{sid}").get_json()["report_rows"][1]["include"])
        self.client.put(f"/api/samples/{sid}/report-print", json={"excludes": []})
        audit_count = sum(1 for event in self.client.get("/api/audit").get_json()
                          if event["action"] == "result_override")
        repeated = self.client.put(f"/api/reports/{sid}/manual", json={
            "rows": rows, "reason": "重复保存检查",
        })
        self.assertFalse(repeated.get_json()["changed"])
        self.assertEqual(audit_count, sum(1 for event in self.client.get("/api/audit").get_json()
                                         if event["action"] == "result_override"))

        wb = load_workbook(io.BytesIO(self.download(f"/api/excel/reports/{sid}")), data_only=True)
        self.assertEqual(12.34, wb["分析报告"]["C13"].value)
        self.assertEqual("补充说明项", wb["分析报告"]["B14"].value)
        wb.close()

        restored = self.client.delete(f"/api/reports/{sid}/manual", json={
            "reason": "恢复系统计算",
        })
        self.assertTrue(restored.get_json()["changed"])
        self.assertIsNone(self.client.get(f"/api/report/{sid}").get_json()["manual_report"])

    def test_result_order_template_and_plain_report_export(self):
        sample_ids = []
        for name, raw in (("矩阵样品甲", 10), ("矩阵样品乙", 20)):
            sid = self.create_sample(name)
            task = self.client.get(f"/api/samples/{sid}").get_json()["items"][0]
            reading = self.client.post("/api/readings", json={
                "sample_analyte_id": task["id"],
            }).get_json()["id"]
            self.client.put(f"/api/readings/{reading}", json={
                "raw": raw, "extra": {}, "use_avg": True,
            })
            sample_ids.append(sid)
        template = self.client.post("/api/result-order-templates", json={
            "name": "银优先", "items": ["Ag", "Cu", "Ag"],
        })
        self.assertEqual(200, template.status_code, template.get_data(as_text=True))
        stored = next(item for item in self.client.get("/api/meta").get_json()["result_order_templates"]
                      if item["id"] == template.get_json()["id"])
        self.assertEqual(["Ag", "Cu"], stored["items"])
        blob = self.download(
            f"/api/excel/results-report?sample_ids={','.join(map(str, sample_ids))}"
            f"&template_id={stored['id']}")
        workbook = load_workbook(io.BytesIO(blob), data_only=True)
        self.assertEqual(["结果报告"], workbook.sheetnames)
        sheet = workbook["结果报告"]
        self.assertEqual("结果报告", sheet["A1"].value)
        self.assertTrue(str(sheet["A2"].value).startswith("日期："))
        self.assertEqual(["来样序号", "样品名称", "Ag", "Cu"],
                         [sheet.cell(3, column).value for column in range(1, 5)])
        self.assertEqual(["矩阵样品甲", "矩阵样品乙"],
                         [sheet.cell(row, 1).value for row in (4, 5)])
        self.assertEqual(["业务测试", "业务测试"],
                         [sheet.cell(row, 2).value for row in (4, 5)])
        expected = [self.client.get(f"/api/report/{sid}").get_json()["groups"][0]["final"]["value"]
                    for sid in sample_ids]
        self.assertEqual(expected, [sheet.cell(row, 3).value for row in (4, 5)])
        self.assertIsNone(sheet["D4"].value)
        self.assertFalse(sheet.tables)
        self.assertTrue(sheet.sheet_view.showGridLines)
        workbook.close()

    def test_special_raw_data_is_recalculated(self):
        method = next(m for m in self.meta["special_methods"] if m["code"] == "moisture")
        created = self.client.post("/api/samples", json={
            "name": "水分专项", "workflow_type": "special", "special_method_id": method["id"],
        })
        sid = created.get_json()["id"]
        self.client.put(f"/api/samples/{sid}/status", json={"status": "queued"})
        self.client.put(f"/api/samples/{sid}/status", json={"status": "measuring"})
        saved = self.client.put(f"/api/special-results/{sid}", json={"raw_data": {"note": "旧工作簿备注"}})
        self.assertEqual(200, saved.status_code, saved.get_data(as_text=True))
        blob = self.download(f"/api/excel/samples/{sid}/data")
        wb = load_workbook(io.BytesIO(blob))
        values = {"pan_weight": 10, "sample_weight": 2, "dry_total": 11.5}
        for row in wb["专项数据"].iter_rows(min_row=4):
            if row[0].value in values:
                row[4].value = values[row[0].value]
        for row in range(wb["专项数据"].max_row, 3, -1):
            if wb["专项数据"].cell(row, 1).value == "note":
                wb["专项数据"].delete_rows(row)
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        response = self.client.post(f"/api/excel/samples/{sid}/data",
                                    data={"file": (io.BytesIO(changed.getvalue()), "special.xlsx")},
                                    content_type="multipart/form-data")
        self.assertEqual(200, response.status_code, response.get_data(as_text=True))
        special = self.client.get(f"/api/samples/{sid}").get_json()["special"]
        self.assertEqual(25, special["calculated_data"]["moisture"])
        self.assertEqual("旧工作簿备注", special["raw_data"]["note"])
        self.assertEqual("completed", special["status"])

    def test_xrf_instrument_task_cannot_be_added_to_data_import(self):
        sid = self.create_sample(xrf=True)
        detail = self.client.get(f"/api/samples/{sid}").get_json()
        aid = detail["items"][0]["analyte_id"]
        xrf_instrument = next(i for i in self.meta["instruments"] if i["itype"] == "xrf")
        db = self.connect()
        try:
            prep_id = db.execute(
                "INSERT INTO preparations(sample_id,name) VALUES(%s,%s) RETURNING id",
                (sid, "XRF只读任务")).fetchone()[0]
            xrf_task_id = db.execute("""INSERT INTO sample_analytes(sample_id,preparation_id,analyte_id,instrument_id)
                          VALUES(%s,%s,%s,%s) RETURNING id""",
                          (sid, prep_id, aid, xrf_instrument["id"])).fetchone()[0]
            db.commit()
        finally:
            db.close()
        blob = self.download(f"/api/excel/samples/{sid}/data")
        wb = load_workbook(io.BytesIO(blob))
        ws = wb["数据录入"]
        ws.append([xrf_task_id, "原样", "Ag", xrf_instrument["name"], "", "", 1, 9.9, 1, 0, "{}", "{}", ""])
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        response = self.client.post(f"/api/excel/samples/{sid}/data",
                                    data={"file": (io.BytesIO(changed.getvalue()), "xrf-task.xlsx")},
                                    content_type="multipart/form-data")
        self.assertEqual(400, response.status_code)
        self.assertIn("XRF任务", response.get_json()["error"])

    def test_write_routes_declare_required_capabilities(self):
        create = lims.app.view_functions["excel_sample_create"]
        plan = lims.app.view_functions["excel_sample_plan_overwrite"]
        data = lims.app.view_functions["excel_sample_data_overwrite"]
        self.assertEqual("sample_manage", inspect.getclosurevars(create).nonlocals["capability"])
        self.assertEqual("sample_manage", inspect.getclosurevars(plan).nonlocals["capability"])
        self.assertEqual("result_edit", inspect.getclosurevars(data).nonlocals["capability"])

    def test_old_backup_routes_are_absent(self):
        self.assertEqual(404, self.client.get("/api/excel/export").status_code)
        self.assertEqual(404, self.client.post("/api/excel/preview").status_code)
        self.assertEqual(404, self.client.post("/api/excel/import").status_code)
        self.assertNotIn("backup_restore", self.client.get("/api/meta").get_json()["capabilities"])


if __name__ == "__main__":
    unittest.main()
