# -*- coding: utf-8 -*-
"""Single-sheet result report export."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _set_value(cell, value):
    cell.value = value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        cell.quotePrefix = True


def build_result_report(columns, sample_rows, report_date):
    """Create a plain sample-by-result grid with no business-workbook metadata."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果报告"
    sheet.sheet_view.showGridLines = True
    headers = ["来样序号", "样品名称", *columns]
    end_column = max(2, len(headers))

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet.cell(1, 1, "结果报告")
    sheet.cell(1, 1).font = Font(name="宋体", size=16, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    sheet.cell(2, 1, f"日期：{report_date}")
    sheet.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column_index, header in enumerate(headers, 1):
        cell = sheet.cell(3, column_index)
        _set_value(cell, header)
        cell.font = Font(name="宋体", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_index, sample in enumerate(sample_rows, 4):
        values = [sample.get("sample_no"), sample.get("sample_name"), *(
            sample.get("values", {}).get(item) for item in columns)]
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column_index)
            _set_value(cell, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 22
    for column_index in range(3, end_column + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 12
    sheet.freeze_panes = "C4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
